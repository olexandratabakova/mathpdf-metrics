"""
CLI entry point using Typer.
Commands: run, compare

New flags (Change 1):
  --dry-run       Print cost estimate and exit without calling the API.
  --max-cost N    Prompt for confirmation if estimated cost > N USD (default from config).
  --no-open       Do not open report.html in the browser after the run.
"""
from __future__ import annotations

import asyncio
import io
import json
import sys
import webbrowser
from pathlib import Path
from typing import Optional

# ── Windows UTF-8 fix ─────────────────────────────────────────────────────────
if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass
if hasattr(sys.stderr, "reconfigure"):
    try:
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

import typer
from rich.console import Console
from rich.table import Table

from . import config  # triggers .env load
from .cost_estimator import CostEstimate, estimate_cost, format_estimate
from .evaluator import evaluate_card_set
from .excel_writer import write_output
from .llm_client import get_token_stats
from .loaders import (
    load_all_card_pdfs,
    load_notes_text,
    match_excel_row_to_cardset,
    pages_to_text_with_markers,
)
from .models import SetMetrics
from .report_writer import write_report

app = typer.Typer(
    name="evaluate",
    help="Evaluate LLM-generated flashcards against source study notes.",
    add_completion=False,
)
console = Console()

_ROOT = Path(__file__).parent.parent
_DEFAULT_NOTES     = _ROOT / "notes" / "main (11).pdf"
_DEFAULT_CARDS_DIR = _ROOT / "cards"
_DEFAULT_EXCEL     = _ROOT / "metrics" / "evaluation MathPDF_all lecture notes.xlsx"
_DEFAULT_OUTPUT_DIR = _ROOT / "output"


def _load_excel_rows(excel_path: Path) -> list[dict]:
    """Load set rows from the Excel file (skip header rows 1-2)."""
    import openpyxl
    wb = openpyxl.load_workbook(str(excel_path))
    ws = wb.active
    rows = []
    for row_idx in range(3, ws.max_row + 1):
        topic = ws.cell(row_idx, 1).value
        card_type = ws.cell(row_idx, 2).value
        level = ws.cell(row_idx, 3).value
        if topic and card_type:
            rows.append({
                "topic": str(topic).strip(),
                "card_type": str(card_type).strip(),
                "level": str(level).strip() if level else None,
                "row_idx": row_idx,
            })
    return rows


def _print_summary(set_results: list[SetMetrics]) -> None:
    """Print a Rich summary table to the terminal."""
    table = Table(title="Evaluation Summary", show_lines=True)
    table.add_column("Topic", style="cyan", no_wrap=True)
    table.add_column("Level", style="magenta")
    table.add_column("N", justify="right")
    table.add_column("Self-Cont%", justify="right")
    table.add_column("Faith%", justify="right")
    table.add_column("Atomic%", justify="right")
    table.add_column("Trivial%", justify="right")
    table.add_column("Dup%", justify="right")
    table.add_column("Coverage%", justify="right")
    table.add_column("MatType")

    for sm in set_results:
        table.add_row(
            sm.topic,
            sm.level or "—",
            str(len(sm.card_metrics)),
            f"{sm.pct_self_contained*100:.0f}%",
            f"{sm.avg_faithfulness*100:.0f}%",
            f"{sm.pct_atomic*100:.0f}%",
            f"{sm.pct_trivial*100:.0f}%",
            f"{sm.duplicate_rate*100:.0f}%",
            f"{sm.coverage_ratio*100:.0f}%",
            sm.material_type or "—",
        )

    console.print(table)

    stats = get_token_stats()
    console.print(
        f"\n[bold]Model:[/bold] {config.JUDGE_MODEL}\n"
        f"[bold]Tokens:[/bold] {stats['prompt_tokens']:,} prompt + "
        f"{stats['completion_tokens']:,} completion = {stats['total_tokens']:,} total\n"
        f"[bold]Actual cost:[/bold] ${stats['estimated_cost_usd']:.4f} USD\n"
    )


def _resolve_card_sets(
    notes: Path,
    cards_dir: Path,
    excel: Path,
    limit: int | None,
    filter_topics: set | None,
) -> list[tuple]:
    """Return [(CardSet, pages_dict)] for all matched/filtered rows."""
    card_sets = load_all_card_pdfs(cards_dir)
    console.print(f"  Found {len(card_sets)} card sets.")

    excel_rows = _load_excel_rows(excel)
    result = []
    for row in excel_rows:
        if filter_topics and row["topic"].lower() not in filter_topics:
            continue
        cs = match_excel_row_to_cardset(
            row["topic"], row["card_type"], row["level"], card_sets
        )
        if cs is None:
            console.print(
                f"  [yellow]No card file found for:[/yellow] "
                f"{row['topic']} / {row['card_type']} / {row['level']}"
            )
            continue
        ps, pe = cs.page_start, cs.page_end
        pages = load_notes_text(notes, ps, pe)
        if not pages:
            console.print(f"  [red]No notes pages found for pp.{ps}-{pe}[/red]")
            continue
        console.print(
            f"  Queued: [cyan]{cs.topic}[/cyan] / {cs.level or '(no level)'} "
            f"| {len(cs.cards[:limit] if limit else cs.cards)} cards | notes pp.{ps}-{pe}"
        )
        result.append((cs, pages))
    return result


@app.command("run")
def cmd_run(
    notes:      Path = typer.Option(_DEFAULT_NOTES,      "--notes",      help="Source notes PDF"),
    cards_dir:  Path = typer.Option(_DEFAULT_CARDS_DIR,  "--cards-dir",  help="Card PDFs directory"),
    excel:      Path = typer.Option(_DEFAULT_EXCEL,       "--excel",      help="Metrics Excel file"),
    output_dir: Path = typer.Option(_DEFAULT_OUTPUT_DIR, "--output-dir", help="Output directory"),
    limit:      Optional[int]  = typer.Option(None,   "--limit", "-n", help="Evaluate first N cards per set"),
    no_cache:   bool = typer.Option(False,  "--no-cache",  help="Disable LLM disk cache"),
    sets_filter:Optional[str]  = typer.Option(None,   "--sets",   help="Comma-sep topics to evaluate"),
    dry_run:    bool = typer.Option(False,  "--dry-run",  help="Estimate cost and exit without calling API"),
    max_cost:   float = typer.Option(None,  "--max-cost", help="Confirm before running if cost > this USD"),
    no_open:    bool = typer.Option(False,  "--no-open",  help="Do not open report.html in browser"),
):
    """Evaluate flashcard sets and fill the metrics spreadsheet + HTML report."""
    config.get_api_key()
    use_cache = not no_cache
    effective_max_cost = max_cost if max_cost is not None else config.MAX_COST

    console.print(f"\n[bold green]Flashcard Evaluator[/bold green]  "
                  f"[dim]model={config.JUDGE_MODEL}[/dim]")
    console.print(f"  Notes:    {notes}")
    console.print(f"  Cards:    {cards_dir}")
    console.print(f"  Excel:    {excel}")
    console.print(f"  Output:   {output_dir}")
    if limit:
        console.print(f"  [yellow]Limit: first {limit} cards per set[/yellow]")
    if no_cache:
        console.print(f"  [yellow]Cache disabled[/yellow]")
    console.print()

    # ── Load sets ─────────────────────────────────────────────────────────────
    console.print("[bold]Loading card PDFs...[/bold]")
    filter_topics = (
        {t.strip().lower() for t in sets_filter.split(",")} if sets_filter else None
    )
    card_sets_with_pages = _resolve_card_sets(notes, cards_dir, excel, limit, filter_topics)

    if not card_sets_with_pages:
        console.print("[red]No sets matched. Check file paths and Excel rows.[/red]")
        raise typer.Exit(1)

    # ── Cost estimation (always printed) ──────────────────────────────────────
    console.print("\n[bold]Estimating cost...[/bold]")
    est = estimate_cost(card_sets_with_pages, limit=limit, model=config.JUDGE_MODEL)
    console.print(format_estimate(est))

    if dry_run:
        console.print("\n[yellow]--dry-run: exiting without calling the API.[/yellow]")
        raise typer.Exit(0)

    # ── Confirm if estimate exceeds max_cost ──────────────────────────────────
    if est.total_cost > effective_max_cost:
        console.print(
            f"\n[bold yellow]⚠  Estimated cost ${est.total_cost:.4f} exceeds "
            f"--max-cost ${effective_max_cost:.2f}.[/bold yellow]"
        )
        confirm = typer.prompt("Type 'yes' to proceed, anything else to abort")
        if confirm.strip().lower() != "yes":
            console.print("[red]Aborted.[/red]")
            raise typer.Exit(1)

    console.print()

    # ── Run evaluation ────────────────────────────────────────────────────────
    details_log: list = []

    async def _run_all() -> list[SetMetrics]:
        results = await asyncio.gather(*[
            evaluate_card_set(cs, pages, limit=limit,
                              use_cache=use_cache, details_log=details_log)
            for cs, pages in card_sets_with_pages
        ])
        return list(results)

    set_results = asyncio.run(_run_all())

    if not set_results:
        console.print("[red]No sets were evaluated.[/red]")
        raise typer.Exit(1)

    # ── Write outputs ─────────────────────────────────────────────────────────
    from datetime import datetime
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_dir.mkdir(parents=True, exist_ok=True)
    details_log_path = output_dir / "details.jsonl"

    excel_out, csv_out = write_output(
        original_excel=excel,
        set_results=set_results,
        output_dir=output_dir,
        details_log_path=details_log_path,
        details_log_data=details_log,
    )

    # ── HTML report ───────────────────────────────────────────────────────────
    report_path = write_report(set_results, output_dir, ts)

    console.print(f"\n[bold green]Done![/bold green]")
    console.print(f"  Excel  → {excel_out}")
    console.print(f"  CSV    → {csv_out}")
    console.print(f"  Log    → {details_log_path}")
    console.print(f"  [bold cyan]Report → {report_path}[/bold cyan]")

    _print_summary(set_results)

    if not no_open:
        console.print(f"\n[dim]Opening report in browser...[/dim]")
        webbrowser.open(report_path.as_uri())


@app.command("compare")
def cmd_compare(
    set_a: Path = typer.Argument(..., help="Card PDF for set A"),
    set_b: Path = typer.Argument(..., help="Card PDF for set B"),
    notes: Path = typer.Option(_DEFAULT_NOTES, "--notes"),
    limit: Optional[int] = typer.Option(None, "--limit", "-n"),
    no_cache: bool = typer.Option(False, "--no-cache"),
    no_open:  bool = typer.Option(False, "--no-open",
                                  help="Do not open report.html in browser"),
):
    """
    Compare two card sets for the same notes (e.g., with/without a user comment).
    Reports differences in card_type distribution and average difficulty.
    """
    from .loaders import load_card_pdf

    config.get_api_key()
    use_cache = not no_cache
    details_log: list = []

    async def _run():
        cs_a = load_card_pdf(set_a)
        cs_b = load_card_pdf(set_b)
        ps = cs_a.page_start or cs_b.page_start
        pe = cs_a.page_end or cs_b.page_end
        pages = load_notes_text(notes, ps, pe)
        sm_a, sm_b = await asyncio.gather(
            evaluate_card_set(cs_a, pages, limit=limit,
                              use_cache=use_cache, details_log=details_log),
            evaluate_card_set(cs_b, pages, limit=limit,
                              use_cache=use_cache, details_log=details_log),
        )
        return sm_a, sm_b

    sm_a, sm_b = asyncio.run(_run())

    console.print("\n[bold]Card Type Distribution:[/bold]")
    table = Table(show_lines=True)
    table.add_column("Card Type")
    table.add_column("Set A %", justify="right")
    table.add_column("Set B %", justify="right")
    table.add_column("Δ", justify="right")

    all_types = set(sm_a.card_type_distribution) | set(sm_b.card_type_distribution)
    for ct in sorted(all_types):
        va = sm_a.card_type_distribution.get(ct, 0.0)
        vb = sm_b.card_type_distribution.get(ct, 0.0)
        delta = vb - va
        color = "green" if delta > 0 else ("red" if delta < 0 else "white")
        table.add_row(ct, f"{va*100:.1f}%", f"{vb*100:.1f}%",
                      f"[{color}]{delta*100:+.1f}%[/{color}]")
    console.print(table)

    avg_diff_a = (
        sum(m.difficulty for m in sm_a.card_metrics if m.difficulty) /
        max(1, sum(1 for m in sm_a.card_metrics if m.difficulty))
    )
    avg_diff_b = (
        sum(m.difficulty for m in sm_b.card_metrics if m.difficulty) /
        max(1, sum(1 for m in sm_b.card_metrics if m.difficulty))
    )
    console.print(
        f"\n[bold]Average Difficulty:[/bold]  "
        f"Set A = {avg_diff_a:.2f}  |  Set B = {avg_diff_b:.2f}  |  "
        f"Δ = {avg_diff_b - avg_diff_a:+.2f}\n"
    )

    _print_summary([sm_a, sm_b])

    # HTML report for compare
    from datetime import datetime
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    report_path = write_report([sm_a, sm_b], _DEFAULT_OUTPUT_DIR, ts)
    console.print(f"\n  [bold cyan]Report → {report_path}[/bold cyan]")
    if not no_open:
        webbrowser.open(report_path.as_uri())


if __name__ == "__main__":
    app()
